"""Read a stockist spreadsheet, find each shop on the map, and merge it into
the `shops` list in places.js.

Used by the "Upload stockists (Excel)" card in Where to find our products.
The CRM calls `read_rows` -> `build_preview` (nothing saved), shows the result,
then `apply_import` on confirm.

Postcodes are looked up with postcodes.io (free, no key, bulk lookups of up to
100 postcodes per request). Rows without a usable postcode fall back to
OpenStreetMap Nominatim, which allows one request per second, so only a few
of those are attempted per upload to stay inside the serverless time limit.
"""

import csv
import io
import itertools
import json
import re
import time
import urllib.error
import urllib.parse
import urllib.request

POSTCODES_BULK_URL = "https://api.postcodes.io/postcodes"
NOMINATIM_URL = "https://nominatim.openstreetmap.org/search"
USER_AGENT = "EmmaBasicCMS/1.0 (+https://emmabasic.co.uk)"
POSTCODE_BATCH = 100
ADDRESS_LOOKUP_LIMIT = 4
ADDRESS_LOOKUP_GAP = 1.1
MAX_ROWS = 2000
MAX_SALES_LINES = 100000

TEMPLATE_COLUMNS = ["Name", "Address", "Town", "Postcode", "Region", "Website", "Phone"]
TEMPLATE_EXAMPLE = [
    "Example Deli", "12 High Street", "Bristol", "BS1 4DJ", "South West",
    "https://exampledeli.co.uk", "0117 000 0000",
]

# Normalised header text (lowercase, punctuation -> single spaces) -> field.
HEADER_ALIASES = {
    "name": ["name", "store", "store name", "shop", "shop name", "stockist", "stockist name",
             "retailer", "retailer name", "business", "business name", "company",
             "company name", "account name", "outlet", "outlet name",
             "bill to name", "ship to name", "sell to name", "sell to customer name"],
    "address": ["address", "address 1", "address line 1", "address1", "addr", "street",
                "street address", "line 1", "full address", "bill to address", "ship to address"],
    "address2": ["address 2", "address line 2", "address2", "line 2",
                 "bill to address 2", "ship to address 2"],
    "town": ["town", "city", "town city", "city town", "locality", "town or city",
             "bill to city", "ship to city", "sell to city", "bill to town", "ship to town"],
    "postcode": ["postcode", "post code", "postal code", "zip", "zip code", "zipcode",
                 "bill to post code", "bill to postcode", "ship to post code",
                 "ship to postcode", "sell to post code", "sell to postcode"],
    "kind": ["vertical name", "vertical", "customer type", "channel", "business type",
             "segment", "customer group", "trade type"],
    "region": ["region", "county", "area", "state", "province"],
    "url": ["website", "url", "link", "web", "site", "web address", "website url",
            "web site", "map link"],
    "phone": ["phone", "telephone", "tel", "phone number", "telephone number",
              "contact number", "mobile"],
    "lat": ["lat", "latitude"],
    "lng": ["lng", "lon", "long", "longitude"],
}
_ALIAS_LOOKUP = {alias: field for field, aliases in HEADER_ALIASES.items() for alias in aliases}

POSTCODE_RE = re.compile(r"\b([A-Z]{1,2}\d[A-Z\d]?)\s*(\d[A-Z]{2})\b", re.IGNORECASE)
_FULL_POSTCODE_RE = re.compile(r"^[A-Z]{1,2}\d[A-Z\d]?\d[A-Z]{2}$")

# Optional shop fields that a spreadsheet can set; blank cells never wipe them.
# Nothing else from the file (prices, quantities, dates, codes) is ever kept.
SHOP_TEXT_FIELDS = ("address", "city", "postcode", "region", "url", "phone")

# Distributor sales exports repeat the customer on every sale line.
_SALES_HEADER_RE = re.compile(r"^(bill|ship|sell) to ")
_STREET_WORDS = {"road", "rd", "street", "st", "lane", "ln", "square", "avenue", "ave",
                 "bridge", "way", "close", "drive", "place", "terrace", "crescent", "parade",
                 "row", "walk", "mews", "yard", "court", "gardens"}
_TOWN_SMALL_WORDS = {"on", "upon", "by", "the", "le", "in", "under", "de", "next", "super", "cum", "of"}

# (vertical / customer type, reason). Matched against the file's customer-type column.
_FLAG_KINDS = [
    ("internet", "Sells online — not a shop people can visit."),
    ("online", "Sells online — not a shop people can visit."),
    ("staff", "Staff purchase or private person — not a shop."),
    ("food service", "Listed as food service (café, restaurant or caterer) — may use the products rather than sell them."),
    ("foodservice", "Listed as food service (café, restaurant or caterer) — may use the products rather than sell them."),
    ("wholesale", "Wholesaler — not a shop."),
    ("distribut", "Distributor — not a shop."),
]
_FLAG_NAME_WORDS = [
    (r"\bclf\b", "This is the distributor itself — not a shop."),
    (r"distribut", "Looks like a distributor — not a shop."),
    (r"wholesal", "Looks like a wholesaler — not a shop."),
    (r"cash\s*(&|and)\s*carry", "Looks like a cash & carry — not a shop."),
    (r"\bimporters?\b", "Looks like an importer or distributor — not a shop."),
    (r"head\s*office", "Looks like a head office — not a shop."),
    (r"^(mr|mrs|ms|miss|dr)\.?\s", "Looks like a private person — not a shop."),
]


class StockistFileError(ValueError):
    """The spreadsheet can't be read. The message is shown to the user."""


class GeocodeError(RuntimeError):
    """The lookup service couldn't be reached."""


# ---------------------------------------------------------------------------
# Small helpers
# ---------------------------------------------------------------------------
def _norm_header(text):
    return re.sub(r"[^a-z0-9]+", " ", str(text or "").lower()).strip()


def _norm_key(text):
    return re.sub(r"[^a-z0-9]+", "", str(text or "").lower())


def _cell_text(value):
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    return re.sub(r"\s+", " ", str(value)).strip()


def normalise_postcode(text):
    """'sw1x7xl' -> 'SW1X 7XL'. Returns '' when it doesn't look like a UK postcode."""
    compact = re.sub(r"\s+", "", str(text or "")).upper()
    if not _FULL_POSTCODE_RE.match(compact):
        return ""
    return compact[:-3] + " " + compact[-3:]


def extract_postcode(text):
    match = POSTCODE_RE.search(str(text or ""))
    return normalise_postcode(match.group(1) + match.group(2)) if match else ""


def _parse_coord(value, low, high):
    try:
        num = float(str(value).strip())
    except (TypeError, ValueError):
        return None
    return num if low <= num <= high else None


def _num(value):
    value = round(float(value), 6)
    return int(value) if value.is_integer() else value


def shop_postcode(shop):
    return normalise_postcode(shop.get("postcode")) or extract_postcode(shop.get("address"))


def shop_key(shop):
    """Shops are the same when name and postcode match (or name and town when
    there is no postcode)."""
    postcode = shop_postcode(shop).replace(" ", "")
    if postcode:
        return (_norm_key(shop.get("name")), postcode)
    return (_norm_key(shop.get("name")), "", _norm_key(shop.get("city")))


def _has_coords(shop):
    return isinstance(shop.get("lat"), (int, float)) and isinstance(shop.get("lng"), (int, float))


def _cap_word(word):
    return "-".join(p[:1].upper() + p[1:].lower() if p else p for p in word.split("-"))


_CONTACT_NOTE_RE = re.compile(r"\(([^)]*\d{5,}[^)]*)\)")


def contact_note_words(names):
    """Words that appear in bracketed notes next to a phone number, e.g. the
    rep name in 'Caner Supermarket (Baki 07828737806)'."""
    words = set()
    for name in names:
        for note in _CONTACT_NOTE_RE.findall(_cell_text(name)):
            words.update(w.lower() for w in re.findall(r"[A-Za-z]+", note))
    return words


def tidy_name(raw, note_words=()):
    """Turn an accounts-system customer name into a shop name for the public
    site. Only removes or re-cases what is there; never adds words.

    'J.S Health Ltd ta Natural Health'      -> 'Natural Health'
    'Caner Supermarket (Baki 07828737806)'  -> 'Caner Supermarket'
    'KRYSTALS EXPRESS LYMINGTON LTD'        -> 'Krystals Express Lymington'

    `note_words` (from contact_note_words) also removes brackets holding only
    those words, e.g. '(Baki)' once 'Baki 07828737806' has been seen.
    """
    name = _cell_text(raw)
    # Notes in brackets that carry a phone number or account code.
    name = re.sub(r"\s*\([^)]*\d{5,}[^)]*\)", "", name)
    if note_words:
        def drop_note(match):
            words = re.findall(r"[A-Za-z]+", match.group(1))
            return "" if words and all(w.lower() in note_words for w in words) else match.group(0)
        name = re.sub(r"\s*\(([^)]*)\)", drop_note, name)
    name = re.sub(r"\s+(t/a|ta|trading as)\s*$", "", name, flags=re.IGNORECASE)
    parts = re.split(r"\s+(?:t/a|ta|trading as)\s+", name, maxsplit=1, flags=re.IGNORECASE)
    if len(parts) == 2 and parts[1].strip():
        name = parts[1]
    legal = r"(ltd\.?|limited|plc|llp)"
    name = re.sub(r"\s+" + legal + r"(?=\s*\()", "", name, flags=re.IGNORECASE)
    while True:
        trimmed = re.sub(r"[\s,]+" + legal + r"\s*$", "", name, flags=re.IGNORECASE)
        if trimmed == name:
            break
        name = trimmed
    name = re.sub(r"\s+", " ", name).strip(" -,.") or _cell_text(raw)
    letters = re.sub(r"[^A-Za-z]", "", name)
    if letters and letters.isupper() and len(letters) > 3:
        words = []
        for word in name.split(" "):
            alpha = re.sub(r"[^A-Za-z]", "", word)
            if len(alpha) <= 2 or (len(alpha) <= 3 and not re.search(r"[AEIOUY]", alpha)):
                words.append(word)
            else:
                words.append(_cap_word(word))
        name = " ".join(words)
    return name


def tidy_town(raw):
    """'BEDFORD' -> 'Bedford', 'Leigh-On-sea' -> 'Leigh-on-Sea'."""
    town = _cell_text(raw)
    if not town or re.search(r"\d", town):
        return town
    out = []
    for i, word in enumerate(town.split(" ")):
        pieces = []
        for j, piece in enumerate(word.split("-")):
            first = i == 0 and j == 0
            if piece.lower() in _TOWN_SMALL_WORDS and not first:
                pieces.append(piece.lower())
            else:
                pieces.append(piece[:1].upper() + piece[1:].lower())
        out.append("-".join(pieces))
    return " ".join(out)


def looks_like_street(text):
    text = _cell_text(text)
    if not text:
        return False
    if re.match(r"^\d", text):
        return True
    last = re.sub(r"[^a-z]", "", text.lower().split(" ")[-1])
    return " " in text and last in _STREET_WORDS


def flag_reason(row):
    """Why this customer is probably not a shop the public can visit ('' if fine)."""
    kind = (row.get("kind") or "").lower()
    for word, reason in _FLAG_KINDS:
        if word in kind:
            return reason
    name = (row.get("name") or "").lower()
    for pattern, reason in _FLAG_NAME_WORDS:
        if re.search(pattern, name):
            return reason
    return ""


# ---------------------------------------------------------------------------
# Reading the spreadsheet
# ---------------------------------------------------------------------------
def match_headers(cells):
    """Map known fields to column indexes. First matching column wins."""
    cols = {}
    for idx, cell in enumerate(cells):
        field = _ALIAS_LOOKUP.get(_norm_header(cell))
        if field and field not in cols:
            cols[field] = idx
    return cols


def _table_from_xlsx(data):
    try:
        import openpyxl
    except ImportError:  # pragma: no cover - listed in requirements.txt
        raise StockistFileError("Excel support isn't installed on the server.")
    try:
        wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except Exception:
        raise StockistFileError(
            "That file couldn't be opened as an Excel workbook. "
            "Please save it as .xlsx (Excel Workbook) and try again."
        )
    try:
        # Pick the sheet whose heading row matches the most known columns, so a
        # pivot/summary sheet in front of the real data is skipped.
        best, best_score = None, 0
        for ws in wb.worksheets:
            for row in itertools.islice(ws.iter_rows(values_only=True), 15):
                found = match_headers(row or ())
                if "name" in found and len(found) > best_score:
                    best, best_score = ws, len(found)
        if best is None:
            best = wb.worksheets[0] if wb.worksheets else None
        return [list(r) for r in best.iter_rows(values_only=True)] if best is not None else []
    finally:
        wb.close()


def _table_from_csv(data):
    for encoding in ("utf-8-sig", "cp1252"):
        try:
            text = data.decode(encoding)
            break
        except UnicodeDecodeError:
            continue
    else:
        raise StockistFileError("That CSV file couldn't be read. Please save it as UTF-8 CSV or .xlsx.")
    return [row for row in csv.reader(io.StringIO(text))]


def read_rows(filename, data):
    """Return one dict per shop row. Raises StockistFileError with a message
    meant for the person uploading."""
    return read_file(filename, data)["rows"]


def read_file(filename, data):
    """Read a stockist list or a distributor sales export.

    Returns {"rows": [...], "lines_read": int, "consolidated": bool}. Sales
    exports (Bill-to Name / Bill-to City / Bill-to Post Code columns) have one
    line per sale; those are merged into one row per shop (name + postcode),
    and names are tidied for the public site. Only the shop columns are read;
    every other column in the file is ignored.
    """
    ext = (filename or "").lower().rsplit(".", 1)[-1] if "." in (filename or "") else ""
    if ext == "xlsx":
        table = _table_from_xlsx(data)
    elif ext == "csv":
        table = _table_from_csv(data)
    else:
        raise StockistFileError(
            "Please upload an Excel file saved as .xlsx (or a .csv file). "
            "In Excel use File → Save As → Excel Workbook (.xlsx)."
        )

    header_idx, cols = None, {}
    for idx, row in enumerate(table[:15]):
        found = match_headers(row)
        if "name" in found:
            header_idx, cols = idx, found
            break
    if header_idx is None:
        raise StockistFileError(
            "We couldn't find the shop name column. The first row of the sheet should have "
            "headings such as Name, Address, Town, Postcode — download the Excel template "
            "to see the layout."
        )

    sales = any(_SALES_HEADER_RE.match(_norm_header(c)) for c in table[header_idx])
    note_words = set()
    if sales:
        name_idx = cols["name"]
        note_words = contact_note_words(
            r[name_idx] for r in table[header_idx + 1:] if r and name_idx < len(r)
        )
    rows, by_key, lines = [], {}, 0
    for offset, raw in enumerate(table[header_idx + 1:], start=header_idx + 2):
        cells = [_cell_text(c) for c in raw]
        if not any(cells):
            continue
        lines += 1
        if lines > MAX_SALES_LINES:
            raise StockistFileError(
                "That file has more than %d lines. Please split it into smaller files." % MAX_SALES_LINES
            )

        def get(field):
            idx = cols.get(field)
            return cells[idx] if idx is not None and idx < len(cells) else ""

        address = ", ".join(p for p in (get("address"), get("address2")) if p)
        postcode = normalise_postcode(get("postcode")) or extract_postcode(address)
        row = {
            "row": offset,
            "name": get("name"),
            "address": address,
            "town": get("town"),
            "postcode": postcode,
            "postcode_raw": get("postcode"),
            "region": get("region"),
            "url": get("url"),
            "phone": get("phone"),
            "lat": _parse_coord(get("lat"), -90, 90),
            "lng": _parse_coord(get("lng"), -180, 180),
            "kind": get("kind"),
        }
        if sales:
            row["name"] = tidy_name(row["name"], note_words)
            row["town"] = tidy_town(row["town"])
            if not row["name"]:
                continue
            key = (_norm_key(row["name"]), postcode.replace(" ", "") or "~" + _norm_key(row["postcode_raw"] or row["town"]))
            first = by_key.get(key)
            if first is not None:
                for field, value in row.items():
                    if value not in ("", None) and first.get(field) in ("", None):
                        first[field] = value
                continue
            by_key[key] = row
        rows.append(row)
        if len(rows) > MAX_ROWS:
            raise StockistFileError(
                "That file has more than %d shops. Please split it into smaller files." % MAX_ROWS
            )
    return {"rows": rows, "lines_read": lines, "consolidated": sales}


# ---------------------------------------------------------------------------
# Looking places up on the map
# ---------------------------------------------------------------------------
def _post_json(url, payload, timeout):
    req = urllib.request.Request(
        url, data=json.dumps(payload).encode("utf-8"), method="POST",
        headers={"Content-Type": "application/json", "User-Agent": USER_AGENT},
    )
    with urllib.request.urlopen(req, timeout=timeout) as resp:
        return json.loads(resp.read().decode("utf-8"))


def _get_json(url, timeout):
    req = urllib.request.Request(url, headers={"User-Agent": USER_AGENT, "Accept-Language": "en-GB"})
    with urllib.request.urlopen(req, timeout=timeout) as resp:
        return json.loads(resp.read().decode("utf-8"))


def _place_name(result):
    parish = result.get("parish") or ""
    if parish and "unparished" not in parish.lower():
        return parish
    return result.get("admin_district") or ""


def lookup_postcodes(postcodes):
    """{postcode: (lat, lng, place name) or None}. Raises GeocodeError if
    postcodes.io is unreachable."""
    unique = list(dict.fromkeys(p for p in postcodes if p))
    found = {}
    for start in range(0, len(unique), POSTCODE_BATCH):
        batch = unique[start:start + POSTCODE_BATCH]
        try:
            payload = _post_json(POSTCODES_BULK_URL, {"postcodes": batch}, timeout=8)
        except (urllib.error.URLError, OSError, ValueError) as exc:
            raise GeocodeError(str(exc))
        for entry in payload.get("result") or []:
            result = entry.get("result") or None
            found[entry.get("query")] = (
                (result["latitude"], result["longitude"], _place_name(result))
                if result and result.get("latitude") is not None else None
            )
        for pc in batch:
            found.setdefault(pc, None)
    return found


def lookup_address(query):
    """(lat, lng) for a free-text UK address via Nominatim, or None."""
    url = NOMINATIM_URL + "?" + urllib.parse.urlencode({
        "q": query, "format": "json", "limit": 1, "countrycodes": "gb",
    })
    try:
        results = _get_json(url, timeout=6)
    except (urllib.error.URLError, OSError, ValueError) as exc:
        raise GeocodeError(str(exc))
    if not results:
        return None
    return (float(results[0]["lat"]), float(results[0]["lon"]))


# ---------------------------------------------------------------------------
# Preview and apply
# ---------------------------------------------------------------------------
def _compose_address(street, town, postcode):
    lower = street.lower()
    tail = " ".join(p for p in (town, postcode) if p and p.lower() not in lower)
    return ", ".join(p for p in (street, tail) if p)


def _shop_from_row(row):
    street, town = row["address"], row["town"]
    # Accounts exports sometimes put a street in the town column.
    if looks_like_street(town):
        street = ", ".join(p for p in (street, town) if p)
        town = ""
    shop = {"name": row["name"]}
    shop["city"] = town or row["region"]
    shop["address"] = _compose_address(street, town, row["postcode"]) if street else ""
    for field in ("postcode", "region", "url", "phone"):
        shop[field] = row[field]
    return shop


def _merge(existing, incoming):
    """Copy of `existing` updated with the non-blank fields of `incoming`."""
    merged = dict(existing)
    merged["name"] = incoming.get("name") or merged.get("name", "")
    for field in SHOP_TEXT_FIELDS:
        value = incoming.get(field)
        if not value:
            continue
        if field == "postcode" and "postcode" not in merged and shop_postcode(merged) == value:
            continue
        merged[field] = value
    if _has_coords(incoming):
        merged["lat"], merged["lng"] = incoming["lat"], incoming["lng"]
    return merged


def _clean_new(shop):
    out = {"name": shop["name"]}
    for field in SHOP_TEXT_FIELDS:
        if shop.get(field):
            out[field] = shop[field]
    out.setdefault("city", "")
    out.setdefault("address", "")
    if _has_coords(shop):
        out["lat"], out["lng"] = shop["lat"], shop["lng"]
    return out


def build_preview(existing_shops, rows, lookup_postcodes=None, lookup_address=None,
                  address_fallback=True):
    """Work out what an upload would do without saving anything.

    With address_fallback=False (used for sales exports, where every real
    customer has a postcode) rows without a usable postcode are reported
    rather than guessed from the town name.

    Returns {"items": [...], "counts": {...}, "removed_if_replace": [names]}.
    Each item has row, name, status (new/updated/unchanged/failed), reason,
    `shop` — the record that would be saved — plus `flag` (why it is probably
    not a shop, or '') and `include` (ticked by default unless flagged).
    """
    lookup_postcodes = lookup_postcodes or globals()["lookup_postcodes"]
    lookup_address = lookup_address or globals()["lookup_address"]
    existing_by_key = {}
    for shop in existing_shops:
        existing_by_key.setdefault(shop_key(shop), shop)

    items, seen = [], {}
    for row in rows:
        flag = flag_reason(row) if row["name"] else ""
        item = {"row": row["row"], "name": row["name"], "status": "", "reason": "", "shop": None,
                "flag": flag, "include": not flag}
        items.append(item)
        if not row["name"]:
            item.update(status="failed", reason="No shop name in this row.")
            continue
        shop = _shop_from_row(row)
        key = shop_key(shop)
        if key in seen:
            item.update(status="failed",
                        reason="Same shop and postcode as row %d — skipped." % seen[key])
            continue
        seen[key] = row["row"]
        match = existing_by_key.get(key)
        item["match"] = match
        item["shop"] = shop
        if row["lat"] is not None and row["lng"] is not None:
            shop["lat"], shop["lng"] = _num(row["lat"]), _num(row["lng"])
        elif match and _has_coords(match):
            shop["lat"], shop["lng"] = match["lat"], match["lng"]
        else:
            item["needs_lookup"] = True

    need_pc = [it for it in items if it.get("needs_lookup") and it["shop"]["postcode"]]
    pc_results, pc_error = {}, False
    if need_pc:
        try:
            pc_results = lookup_postcodes([it["shop"]["postcode"] for it in need_pc])
        except GeocodeError:
            pc_error = True

    address_budget = ADDRESS_LOOKUP_LIMIT
    last_address_lookup = 0.0
    for row, item in zip(rows, items):
        if not item.get("needs_lookup"):
            continue
        shop = item["shop"]
        coords = pc_results.get(shop["postcode"]) if shop["postcode"] else None
        bad_postcode = bool(row.get("postcode_raw")) and not shop["postcode"]
        if (coords is None and address_fallback and not bad_postcode
                and (row["address"] or row["town"]) and address_budget > 0):
            query = ", ".join(p for p in (row["address"], row["town"], row["postcode"]) if p)
            address_budget -= 1
            wait = ADDRESS_LOOKUP_GAP - (time.monotonic() - last_address_lookup)
            if last_address_lookup and wait > 0:
                time.sleep(wait)
            try:
                coords = lookup_address(query + ", United Kingdom")
            except GeocodeError:
                coords = None
            last_address_lookup = time.monotonic()
        if coords:
            shop["lat"], shop["lng"] = _num(coords[0]), _num(coords[1])
            if not shop["city"] and len(coords) > 2 and coords[2]:
                shop["city"] = coords[2]
            continue
        if pc_error:
            reason = "The postcode lookup service didn't answer — please try the upload again in a minute."
        elif shop["postcode"]:
            reason = "Postcode %s not found — please check it." % shop["postcode"]
        elif row.get("postcode_raw"):
            reason = "“%s” doesn't look like a UK postcode — please check it." % row["postcode_raw"]
        elif not address_fallback:
            reason = "No postcode — we can't place this shop on the map. Please add a UK postcode."
        elif row["address"] or row["town"]:
            reason = ("No postcode, and we couldn't find the address on the map. "
                      "Please add a postcode.")
        else:
            reason = "No postcode or address — we can't place this shop on the map."
        item.update(status="failed", reason=reason)

    counts = {"new": 0, "updated": 0, "unchanged": 0, "failed": 0, "flagged": 0}
    in_file = set()
    for item in items:
        match = item.pop("match", None)
        item.pop("needs_lookup", None)
        if item["status"] == "failed":
            item["shop"] = None
            item["include"] = False
            counts["failed"] += 1
            continue
        in_file.add(shop_key(item["shop"]))
        if match is None:
            item["status"] = "new"
            shop = item["shop"]
            if not shop["address"]:
                shop["address"] = " ".join(p for p in (shop["city"], shop["postcode"]) if p)
            item["shop"] = _clean_new(shop)
        elif _merge(match, item["shop"]) == match:
            item["status"] = "unchanged"
        else:
            item["status"] = "updated"
        counts[item["status"]] += 1
        if item["flag"]:
            counts["flagged"] += 1

    removed = [s.get("name", "") for s in existing_shops if shop_key(s) not in in_file]
    return {"items": items, "counts": counts, "removed_if_replace": removed}


def clean_payload_shop(raw):
    """Validate a shop that came back from the preview page. Returns None if unusable."""
    if not isinstance(raw, dict):
        return None
    name = _cell_text(raw.get("name"))
    if not name:
        return None
    shop = {"name": name}
    for field in SHOP_TEXT_FIELDS:
        value = _cell_text(raw.get(field))
        if value:
            shop[field] = value[:500]
    lat = _parse_coord(raw.get("lat"), -90, 90)
    lng = _parse_coord(raw.get("lng"), -180, 180)
    if lat is None or lng is None:
        return None
    shop["lat"], shop["lng"] = _num(lat), _num(lng)
    return shop


def apply_import(existing_shops, shops, mode):
    """Return (new shop list, summary). mode is "merge" or "replace"."""
    by_key = {}
    for idx, shop in enumerate(existing_shops):
        by_key.setdefault(shop_key(shop), idx)
    summary = {"new": 0, "updated": 0, "unchanged": 0, "removed": 0}

    if mode == "replace":
        result, kept = [], set()
        for shop in shops:
            idx = by_key.get(shop_key(shop))
            if idx is None:
                result.append(_clean_new(shop))
                summary["new"] += 1
                continue
            kept.add(idx)
            merged = _merge(existing_shops[idx], shop)
            summary["unchanged" if merged == existing_shops[idx] else "updated"] += 1
            result.append(merged)
        summary["removed"] = len(existing_shops) - len(kept)
        return result, summary

    result = [dict(s) for s in existing_shops]
    for shop in shops:
        key = shop_key(shop)
        idx = by_key.get(key)
        if idx is None:
            result.append(_clean_new(shop))
            by_key[key] = len(result) - 1
            summary["new"] += 1
            continue
        merged = _merge(result[idx], shop)
        summary["unchanged" if merged == result[idx] else "updated"] += 1
        result[idx] = merged
    return result, summary


def fill_missing_coordinates(shops, lookup_postcodes=None):
    """Give shops without a pin coordinates from their postcode (in place).
    Lookup failures are ignored so saving never depends on the network."""
    lookup_postcodes = lookup_postcodes or globals()["lookup_postcodes"]
    todo = [(s, shop_postcode(s)) for s in shops if not _has_coords(s)]
    todo = [(s, pc) for s, pc in todo if pc]
    if not todo:
        return 0
    try:
        found = lookup_postcodes([pc for _, pc in todo])
    except GeocodeError:
        return 0
    filled = 0
    for shop, pc in todo:
        coords = found.get(pc)
        if coords:
            shop["lat"], shop["lng"] = _num(coords[0]), _num(coords[1])
            filled += 1
    return filled


def template_workbook_bytes():
    import openpyxl
    from openpyxl.styles import Font

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Stockists"
    ws.append(TEMPLATE_COLUMNS)
    ws.append(TEMPLATE_EXAMPLE)
    for cell in ws[1]:
        cell.font = Font(bold=True)
    for letter, width in zip("ABCDEFG", (28, 32, 18, 12, 18, 34, 18)):
        ws.column_dimensions[letter].width = width
    ws.freeze_panes = "A2"
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
