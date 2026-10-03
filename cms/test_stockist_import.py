"""Stockists can be uploaded from a spreadsheet in Where to find our products.

The upload is a two-step flow: the file is read and every row is looked up on
the map, the CRM shows a preview, and only the confirm step writes places.js.
Postcode lookups go over the network, so tests swap in fake lookups.

Run with:  python test_stockist_import.py
"""

import csv
import io
import json
import os
import re
import sys
import unittest
from unittest import mock

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))

import openpyxl  # noqa: E402

import stockists_import as si  # noqa: E402
from app import app, load_data, save_data  # noqa: E402


KNOWN = {
    "W1A 1AB": (51.5143, -0.153),
    "SW1X 7XL": (51.4994, -0.1632),
    "BS1 4DJ": (51.4545, -2.5879),
    "EH1 1YZ": (55.9521, -3.1890),
    "M1 1AE": (53.4794, -2.2453),
}


def fake_postcodes(postcodes):
    return {pc: KNOWN.get(pc) for pc in postcodes}


def no_address_lookup(query):
    return None


def xlsx_bytes(rows):
    wb = openpyxl.Workbook()
    ws = wb.active
    for row in rows:
        ws.append(row)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def csv_bytes(rows):
    buf = io.StringIO()
    writer = csv.writer(buf)
    for row in rows:
        writer.writerow(row)
    return buf.getvalue().encode("utf-8-sig")


EXISTING = [
    {"name": "Selfridges", "city": "London", "address": "400 Oxford St, London W1A 1AB",
     "lat": 51.5143, "lng": -0.153},
    {"name": "Harrods", "city": "London", "address": "87–135 Brompton Rd, London SW1X 7XL",
     "lat": 51.4994, "lng": -0.1632},
    {"name": "Quince & Cook", "city": "Scotland", "address": "Perth, Scotland",
     "lat": 56.395, "lng": -3.431},
]


class HeaderMatchingTests(unittest.TestCase):
    def test_common_header_names_are_recognised_whatever_the_case(self):
        cols = si.match_headers(
            ["Store Name", "ADDRESS 1", "Town/City", "Post Code", "County", "Website", "Telephone"]
        )
        self.assertEqual(cols, {
            "name": 0, "address": 1, "town": 2, "postcode": 3,
            "region": 4, "url": 5, "phone": 6,
        })

    def test_alternative_names(self):
        cols = si.match_headers(["Shop", "Street", "City", "Postcode", "Region", "URL", "Phone"])
        self.assertEqual(cols["name"], 0)
        self.assertEqual(cols["address"], 1)
        self.assertEqual(cols["town"], 2)
        self.assertEqual(cols["url"], 5)
        cols = si.match_headers(["Store", "Address", "Postal code", "Link", "Latitude", "Longitude"])
        self.assertEqual(cols["postcode"], 2)
        self.assertEqual(cols["url"], 3)
        self.assertEqual(cols["lat"], 4)
        self.assertEqual(cols["lng"], 5)

    def test_unknown_columns_are_ignored(self):
        cols = si.match_headers(["Name", "Notes", "Postcode"])
        self.assertEqual(cols, {"name": 0, "postcode": 2})


class ReadRowsTests(unittest.TestCase):
    def test_reads_xlsx_skipping_blank_rows_and_a_title_row(self):
        data = xlsx_bytes([
            ["Our stockists 2026"],
            [],
            ["Shop name", "Address", "Town", "Postcode", "Website"],
            ["Harvey Nichols", "27 Philpot Ln", "Bristol", "bs1 4dj", "https://hn.example"],
            [None, None, None, None, None],
            ["  Valvona  ", "19 Elm Row", "Edinburgh", "EH1 1YZ", ""],
        ])
        rows = si.read_rows("stores.xlsx", data)
        self.assertEqual(len(rows), 2)
        self.assertEqual(rows[0]["name"], "Harvey Nichols")
        self.assertEqual(rows[0]["postcode"], "BS1 4DJ")
        self.assertEqual(rows[0]["row"], 4)
        self.assertEqual(rows[1]["name"], "Valvona")
        self.assertEqual(rows[1]["row"], 6)

    def test_reads_csv(self):
        data = csv_bytes([
            ["Store", "Street", "City", "Post code"],
            ["Harvey Nichols", "27 Philpot Ln", "Bristol", "BS14DJ"],
        ])
        rows = si.read_rows("stores.csv", data)
        self.assertEqual(len(rows), 1)
        self.assertEqual(rows[0]["town"], "Bristol")
        self.assertEqual(rows[0]["postcode"], "BS1 4DJ")

    def test_postcode_is_taken_from_the_address_when_there_is_no_postcode_column(self):
        data = csv_bytes([["Name", "Address"], ["Harrods", "87 Brompton Rd, London SW1X 7XL"]])
        rows = si.read_rows("stores.csv", data)
        self.assertEqual(rows[0]["postcode"], "SW1X 7XL")

    def test_missing_name_column_is_a_plain_english_error(self):
        data = csv_bytes([["Address", "Postcode"], ["1 High St", "BS1 4DJ"]])
        with self.assertRaises(si.StockistFileError) as ctx:
            si.read_rows("stores.csv", data)
        self.assertIn("Name", str(ctx.exception))

    def test_unsupported_file_type(self):
        with self.assertRaises(si.StockistFileError) as ctx:
            si.read_rows("stores.xls", b"junk")
        self.assertIn(".xlsx", str(ctx.exception))


class PreviewTests(unittest.TestCase):
    def preview(self, rows, existing=None):
        return si.build_preview(
            EXISTING if existing is None else existing, rows,
            lookup_postcodes=fake_postcodes, lookup_address=no_address_lookup,
        )

    def rows(self, *data):
        return si.read_rows("s.csv", csv_bytes(
            [["Name", "Address", "Town", "Postcode", "Website"]] + list(data)
        ))

    def test_new_updated_unchanged_and_failed_rows(self):
        result = self.preview(self.rows(
            ["Selfridges", "400 Oxford St", "London", "W1A 1AB", ""],
            ["Harrods", "87–135 Brompton Rd", "London", "SW1X 7XL", "https://harrods.example"],
            ["Harvey Nichols", "27 Philpot Ln", "Bristol", "BS1 4DJ", ""],
            ["Nowhere Deli", "1 Fake St", "Atlantis", "ZZ9 9ZZ", ""],
        ))
        statuses = [item["status"] for item in result["items"]]
        self.assertEqual(statuses, ["unchanged", "updated", "new", "failed"])
        self.assertEqual(result["counts"],
                         {"new": 1, "updated": 1, "unchanged": 1, "failed": 1, "flagged": 0})

    def test_failed_postcode_row_explains_why(self):
        result = self.preview(self.rows(["Nowhere Deli", "1 Fake St", "Atlantis", "ZZ9 9ZZ", ""]))
        failed = result["items"][0]
        self.assertEqual(failed["status"], "failed")
        self.assertIn("ZZ9 9ZZ", failed["reason"])
        self.assertIn("not found", failed["reason"].lower())

    def test_row_without_name_fails(self):
        result = self.preview(self.rows(["", "1 High St", "Bristol", "BS1 4DJ", ""]))
        self.assertEqual(result["items"][0]["status"], "failed")
        self.assertIn("name", result["items"][0]["reason"].lower())

    def test_new_store_gets_coordinates_and_a_full_address(self):
        result = self.preview(self.rows(["Harvey Nichols", "27 Philpot Ln", "Bristol", "BS1 4DJ", ""]))
        shop = result["items"][0]["shop"]
        self.assertEqual(shop["lat"], 51.4545)
        self.assertEqual(shop["lng"], -2.5879)
        self.assertEqual(shop["address"], "27 Philpot Ln, Bristol BS1 4DJ")
        self.assertEqual(shop["postcode"], "BS1 4DJ")
        self.assertEqual(shop["city"], "Bristol")

    def test_existing_pin_is_kept_when_postcode_is_unchanged(self):
        calls = []

        def counting(postcodes):
            calls.extend(postcodes)
            return fake_postcodes(postcodes)

        existing = [dict(EXISTING[0], lat=51.51, lng=-0.15)]
        result = si.build_preview(
            existing, self.rows(["Selfridges", "400 Oxford St", "London", "W1A 1AB", ""]),
            lookup_postcodes=counting, lookup_address=no_address_lookup,
        )
        self.assertEqual(calls, [])
        self.assertEqual(result["items"][0]["shop"]["lat"], 51.51)

    def test_explicit_coordinates_in_the_file_are_used(self):
        rows = si.read_rows("s.csv", csv_bytes([
            ["Name", "Postcode", "Lat", "Lng"],
            ["Pop-up", "ZZ9 9ZZ", "52.1", "-1.2"],
        ]))
        result = self.preview(rows)
        self.assertEqual(result["items"][0]["status"], "new")
        self.assertEqual(result["items"][0]["shop"]["lat"], 52.1)

    def test_no_postcode_falls_back_to_an_address_search(self):
        rows = si.read_rows("s.csv", csv_bytes([["Name", "Town"], ["Village Shop", "Hay-on-Wye"]]))
        result = si.build_preview(
            [], rows, lookup_postcodes=fake_postcodes,
            lookup_address=lambda q: (52.07, -3.13) if "Hay-on-Wye" in q else None,
        )
        self.assertEqual(result["items"][0]["status"], "new")
        self.assertEqual(result["items"][0]["shop"]["lng"], -3.13)

    def test_duplicate_rows_in_the_file_are_reported(self):
        result = self.preview(self.rows(
            ["Harvey Nichols", "27 Philpot Ln", "Bristol", "BS1 4DJ", ""],
            ["harvey nichols", "27 Philpot Lane", "Bristol", "BS14DJ", ""],
        ))
        self.assertEqual(result["items"][1]["status"], "failed")
        self.assertIn("row 2", result["items"][1]["reason"].lower())

    def test_preview_lists_shops_a_replace_would_remove(self):
        result = self.preview(self.rows(["Selfridges", "400 Oxford St", "London", "W1A 1AB", ""]))
        self.assertEqual(result["removed_if_replace"], ["Harrods", "Quince & Cook"])


class ApplyTests(unittest.TestCase):
    def ready_shops(self, *data):
        rows = si.read_rows("s.csv", csv_bytes(
            [["Name", "Address", "Town", "Postcode", "Website"]] + list(data)
        ))
        preview = si.build_preview(EXISTING, rows, lookup_postcodes=fake_postcodes,
                                   lookup_address=no_address_lookup)
        return [item["shop"] for item in preview["items"] if item["status"] != "failed"]

    def test_merge_adds_new_updates_matches_and_keeps_the_rest(self):
        shops = self.ready_shops(
            ["Harrods", "", "London", "SW1X 7XL", "https://harrods.example"],
            ["Harvey Nichols", "27 Philpot Ln", "Bristol", "BS1 4DJ", ""],
        )
        merged, summary = si.apply_import(EXISTING, shops, "merge")
        names = [s["name"] for s in merged]
        self.assertEqual(names, ["Selfridges", "Harrods", "Quince & Cook", "Harvey Nichols"])
        harrods = merged[1]
        self.assertEqual(harrods["url"], "https://harrods.example")
        # A blank address cell must not wipe the address already on the site.
        self.assertEqual(harrods["address"], "87–135 Brompton Rd, London SW1X 7XL")
        self.assertEqual(summary, {"new": 1, "updated": 1, "unchanged": 0, "removed": 0})

    def test_replace_keeps_only_the_file(self):
        shops = self.ready_shops(
            ["Harrods", "", "London", "SW1X 7XL", ""],
            ["Harvey Nichols", "27 Philpot Ln", "Bristol", "BS1 4DJ", ""],
        )
        replaced, summary = si.apply_import(EXISTING, shops, "replace")
        self.assertEqual([s["name"] for s in replaced], ["Harrods", "Harvey Nichols"])
        self.assertEqual(replaced[0]["address"], "87–135 Brompton Rd, London SW1X 7XL")
        self.assertEqual(summary["removed"], 2)

    def test_does_not_modify_the_list_passed_in(self):
        before = json.dumps(EXISTING)
        si.apply_import(EXISTING, self.ready_shops(["Harrods", "", "", "SW1X 7XL", "https://x"]), "merge")
        self.assertEqual(json.dumps(EXISTING), before)


SALES_HEADER = ["Sales Date", "Quantity", "CLF_Code", "Trade_Price", "Vendor_Item_Number",
                "Description", "Description 2", "Bill-to Name", "Bill-to City",
                "Bill-to Post Code", "Vertical Name"]


def sales_line(name, city, postcode, vertical="Health Food Retailer", qty=6, price=12.34):
    return ["2026-01-15", qty, "XQ999", price, "INV-77821", "Konjac noodles", "200g",
            name, city, postcode, vertical]


def sales_workbook():
    """Shaped like a distributor sales-history export: a pivot sheet first,
    then one line per sale with the customer repeated."""
    wb = openpyxl.Workbook()
    pivot = wb.active
    pivot.title = "Sheet1"
    pivot.append([None])
    pivot.append(["Row Labels", "Sum of Quantity"])
    pivot.append(["BLUKOO Ltd", 40])
    data = wb.create_sheet("Data 12.12.25-12.03.26")
    data.append(SALES_HEADER)
    for row in [
        sales_line("BLUKOO Ltd", "Kingsbury", "NW9 8AU", "Internet Bulk"),
        sales_line("BLUKOO Ltd", "Kingsbury", "nw98au", "Internet Bulk", qty=3),
        sales_line("J.S Health Ltd ta Natural Health", "Welwyn Garden City", "AL8 6PH"),
        sales_line("J.S Health Ltd ta Natural Health", "Welwyn Garden City", "AL8 6PH", qty=9),
        sales_line("Caner Supermarket (Baki 07828737806)", "London", "E7 9DU", "Convenience"),
        sales_line("KRYSTALS EXPRESS LYMINGTON LTD", "BEDFORD", "MK42 9TW"),
        sales_line("Fairway Importers Ltd", "Chelmsford", "CM1 3SH", "Internet Bulk"),
        sales_line("Fairway Importers Ltd", "Huntingdon", "PE29 1HU", "Internet Bulk"),
        sales_line("Dugard and Daughters (Herne Hill)", "London", "SE24 0EZ"),
        sales_line("Katarzyna Migdal", "Southampton", "SO15 5NF", "Staff"),
        sales_line("The Deli Downstairs", "London", "E9 7JN", "Food Service"),
        sales_line("Underbank General Store", "109 George Leigh Street", "M4 6GG"),
        sales_line("Natures Hand", "Limerick", "V94VW57"),
        sales_line("Greens Health Foods Limited", "Leigh-On-sea", "SS9 2HA"),
        sales_line("Refill Therapy Ltd (Baki)", "London", "E9 5LH", "Convenience"),
        sales_line("Open Sesame Gort", "Gort", None),
    ]:
        data.append(row)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


SALES_KNOWN = {
    "NW9 8AU": (51.5833, -0.2631, "Brent"),
    "AL8 6PH": (51.80, -0.21, "Welwyn Hatfield"),
    "E7 9DU": (51.55, 0.03, "Newham"),
    "MK42 9TW": (52.12, -0.46, "Bedford"),
    "CM1 3SH": (51.74, 0.47, "Chelmsford"),
    "PE29 1HU": (52.33, -0.18, "Huntingdonshire"),
    "SE24 0EZ": (51.45, -0.10, "Lambeth"),
    "SO15 5NF": (50.92, -1.43, "Southampton"),
    "E9 7JN": (51.54, -0.04, "Hackney"),
    "M4 6GG": (53.48, -2.23, "Manchester"),
    "SS9 2HA": (51.54, 0.65, "Southend-on-Sea"),
    "E9 5LH": (51.54, -0.05, "Hackney"),
}


def fake_sales_postcodes(postcodes):
    return {pc: SALES_KNOWN.get(pc) for pc in postcodes}


class SalesExportTests(unittest.TestCase):
    def read(self):
        return si.read_file("CLFSales 12-03-2026.xlsx", sales_workbook())

    def by_postcode(self, rows):
        return {r["postcode"]: r for r in rows}

    def test_bill_to_headers_are_recognised(self):
        cols = si.match_headers(SALES_HEADER)
        self.assertEqual(cols["name"], 7)
        self.assertEqual(cols["town"], 8)
        self.assertEqual(cols["postcode"], 9)
        self.assertEqual(cols["kind"], 10)

    def test_sales_lines_are_consolidated_to_one_row_per_shop(self):
        result = self.read()
        self.assertTrue(result["consolidated"])
        self.assertEqual(result["lines_read"], 16)
        self.assertEqual(len(result["rows"]), 14)

    def test_same_name_with_different_postcodes_is_two_branches(self):
        rows = self.by_postcode(self.read()["rows"])
        self.assertEqual(rows["CM1 3SH"]["name"], "Fairway Importers")
        self.assertEqual(rows["PE29 1HU"]["name"], "Fairway Importers")

    def test_names_are_tidied_for_the_public_site(self):
        rows = self.by_postcode(self.read()["rows"])
        self.assertEqual(rows["NW9 8AU"]["name"], "Blukoo")
        self.assertEqual(rows["AL8 6PH"]["name"], "Natural Health")
        self.assertEqual(rows["E7 9DU"]["name"], "Caner Supermarket")
        self.assertEqual(rows["MK42 9TW"]["name"], "Krystals Express Lymington")
        self.assertEqual(rows["SE24 0EZ"]["name"], "Dugard and Daughters (Herne Hill)")
        self.assertEqual(rows["SS9 2HA"]["name"], "Greens Health Foods")

    def test_rep_names_seen_next_to_phone_numbers_are_removed_everywhere(self):
        rows = self.by_postcode(self.read()["rows"])
        self.assertEqual(rows["E9 5LH"]["name"], "Refill Therapy")
        self.assertNotIn("Baki", json.dumps(self.read()["rows"]))

    def test_towns_are_tidied(self):
        rows = self.by_postcode(self.read()["rows"])
        self.assertEqual(rows["MK42 9TW"]["town"], "Bedford")
        self.assertEqual(rows["SS9 2HA"]["town"], "Leigh-on-Sea")

    def test_no_sales_data_is_kept(self):
        rows = self.read()["rows"]
        dumped = json.dumps(rows)
        for secret in ("12.34", "INV-77821", "XQ999", "2026-01-15", "Konjac", "07828737806", "Quantity"):
            self.assertNotIn(secret, dumped)
        allowed = {"row", "name", "address", "town", "postcode", "postcode_raw", "region",
                   "url", "phone", "lat", "lng", "kind"}
        for row in rows:
            self.assertLessEqual(set(row), allowed)

    def test_pivot_sheet_is_skipped_for_the_data_sheet(self):
        self.assertIn("Natural Health", [r["name"] for r in self.read()["rows"]])

    def preview(self, existing=None):
        return si.build_preview(existing or [], self.read()["rows"],
                                lookup_postcodes=fake_sales_postcodes,
                                lookup_address=no_address_lookup)

    def test_non_retail_customers_are_flagged_and_left_out_by_default(self):
        items = {i["shop"]["postcode"]: i for i in self.preview()["items"] if i["shop"]}
        for pc, word in (("NW9 8AU", "online"), ("CM1 3SH", "online"),
                         ("SO15 5NF", "staff"), ("E9 7JN", "food service")):
            self.assertTrue(items[pc]["flag"], pc)
            self.assertIn(word, items[pc]["flag"].lower(), pc)
            self.assertFalse(items[pc]["include"], pc)
        self.assertEqual(items["AL8 6PH"]["flag"], "")
        self.assertTrue(items["AL8 6PH"]["include"])

    def test_importer_and_wholesaler_names_are_flagged(self):
        reason = si.flag_reason({"name": "Fairway Importers", "kind": ""})
        self.assertIn("importer", reason.lower())
        self.assertIn("wholesale", si.flag_reason({"name": "ABC Wholesale", "kind": ""}).lower())
        self.assertIn("distribut", si.flag_reason({"name": "CLF Distribution", "kind": ""}).lower())
        self.assertEqual(si.flag_reason({"name": "Natural Health", "kind": "Health Food Retailer"}), "")

    def test_counts_include_flagged(self):
        counts = self.preview()["counts"]
        self.assertEqual(counts["failed"], 2)
        self.assertEqual(counts["flagged"], 5)
        self.assertEqual(counts["new"], 12)

    def test_non_uk_postcode_fails_with_a_reason(self):
        failed = {i["name"]: i for i in self.preview()["items"] if i["status"] == "failed"}
        self.assertEqual(sorted(failed), ["Natures Hand", "Open Sesame Gort"])
        self.assertIn("UK postcode", failed["Natures Hand"]["reason"])

    def test_without_address_fallback_a_town_alone_is_not_guessed(self):
        result = si.build_preview(
            [], self.read()["rows"], lookup_postcodes=fake_sales_postcodes,
            lookup_address=lambda q: (55.1, -6.9), address_fallback=False,
        )
        failed = {i["name"]: i for i in result["items"] if i["status"] == "failed"}
        self.assertIn("Open Sesame Gort", failed)
        self.assertIn("postcode", failed["Open Sesame Gort"]["reason"].lower())

    def test_street_in_the_town_column_becomes_the_address(self):
        items = {i["shop"]["postcode"]: i for i in self.preview()["items"] if i["shop"]}
        shop = items["M4 6GG"]["shop"]
        self.assertEqual(shop["city"], "Manchester")
        self.assertIn("109 George Leigh Street", shop["address"])

    def test_new_shop_without_street_gets_town_and_postcode_as_address(self):
        items = {i["shop"]["postcode"]: i for i in self.preview()["items"] if i["shop"]}
        self.assertEqual(items["MK42 9TW"]["shop"]["address"], "Bedford MK42 9TW")

    def test_matches_an_existing_stockist_after_tidying(self):
        existing = [{"name": "Blukoo", "city": "London",
                     "address": "Barningham Way, Kingsbury, London NW9 8AU",
                     "lat": 51.5833, "lng": -0.2631}]
        items = {i["shop"]["postcode"]: i for i in self.preview(existing)["items"] if i["shop"]}
        self.assertIn(items["NW9 8AU"]["status"], ("unchanged", "updated"))


class TemplateTests(unittest.TestCase):
    def test_template_has_the_expected_columns(self):
        wb = openpyxl.load_workbook(io.BytesIO(si.template_workbook_bytes()))
        header = [c.value for c in wb.active[1]]
        self.assertEqual(header, ["Name", "Address", "Town", "Postcode", "Region", "Website", "Phone"])
        rows = si.read_rows("template.xlsx", si.template_workbook_bytes())
        self.assertEqual(len(rows), 1)


class GeocoderTests(unittest.TestCase):
    def test_bulk_lookup_batches_100_postcodes_per_request(self):
        sent = []

        def fake_post(url, payload, timeout):
            sent.append(payload["postcodes"])
            return {"status": 200, "result": [
                {"query": pc, "result": {"latitude": 50.0, "longitude": -1.0,
                                         "parish": "Totnes", "admin_district": "South Hams"}}
                for pc in payload["postcodes"]
            ]}

        postcodes = ["PC%03d" % i for i in range(250)]
        with mock.patch.object(si, "_post_json", fake_post):
            found = si.lookup_postcodes(postcodes)
        self.assertEqual([len(batch) for batch in sent], [100, 100, 50])
        self.assertEqual(len(found), 250)
        self.assertEqual(found[postcodes[0]], (50.0, -1.0, "Totnes"))

    def test_place_name_skips_unparished_areas(self):
        def fake_post(url, payload, timeout):
            return {"status": 200, "result": [{"query": "M4 6GG", "result": {
                "latitude": 53.48, "longitude": -2.23,
                "parish": "Manchester, unparished area", "admin_district": "Manchester"}}]}

        with mock.patch.object(si, "_post_json", fake_post):
            self.assertEqual(si.lookup_postcodes(["M4 6GG"])["M4 6GG"][2], "Manchester")

    def test_unknown_postcode_maps_to_none(self):
        def fake_post(url, payload, timeout):
            return {"status": 200, "result": [{"query": "ZZ9 9ZZ", "result": None}]}

        with mock.patch.object(si, "_post_json", fake_post):
            self.assertEqual(si.lookup_postcodes(["ZZ9 9ZZ"]), {"ZZ9 9ZZ": None})


def _restore(original):
    with app.app_context():
        save_data("places", original)


class ImportRouteTests(unittest.TestCase):
    def setUp(self):
        app.config["TESTING"] = True
        self.client = app.test_client()
        self.patches = [
            mock.patch.object(si, "lookup_postcodes", fake_postcodes),
            mock.patch.object(si, "lookup_address", no_address_lookup),
        ]
        for p in self.patches:
            p.start()

    def tearDown(self):
        for p in self.patches:
            p.stop()

    def upload(self, rows, filename="stores.xlsx"):
        data = xlsx_bytes(rows) if filename.endswith(".xlsx") else csv_bytes(rows)
        return self.client.post(
            "/places/import/preview",
            data={"file": (io.BytesIO(data), filename)},
            content_type="multipart/form-data",
        )

    def test_places_page_has_the_upload_card(self):
        html = self.client.get("/places").get_data(as_text=True)
        self.assertIn("Upload stockists (Excel)", html)
        self.assertIn("/places/import/preview", html)
        self.assertIn("Download Excel template", html)
        self.assertIn("/places/import/template", html)

    def test_template_download(self):
        response = self.client.get("/places/import/template")
        self.assertEqual(response.status_code, 200)
        self.assertIn("attachment", response.headers.get("Content-Disposition", ""))
        self.assertIn(".xlsx", response.headers.get("Content-Disposition", ""))
        wb = openpyxl.load_workbook(io.BytesIO(response.data))
        self.assertEqual(wb.active["A1"].value, "Name")

    def test_preview_shows_counts_and_does_not_save(self):
        before = load_data("places")
        response = self.upload([
            ["Name", "Address", "Town", "Postcode"],
            ["Harvey Nichols", "27 Philpot Ln", "Bristol", "BS1 4DJ"],
            ["Nowhere Deli", "1 Fake St", "Atlantis", "ZZ9 9ZZ"],
        ])
        self.assertEqual(response.status_code, 200)
        html = response.get_data(as_text=True)
        self.assertIn("Harvey Nichols", html)
        self.assertIn("ZZ9 9ZZ", html)
        self.assertRegex(html, r"1(</strong>)?\s+new")
        self.assertRegex(html, r"1(</strong>)?\s+could not be added")
        self.assertIn("/places/import/confirm", html)
        self.assertIn('value="merge" checked', html)
        self.assertEqual(load_data("places"), before)

    def test_preview_of_a_bad_file_explains_the_problem(self):
        response = self.upload([["Colour", "Size"], ["red", "L"]], filename="stores.csv")
        self.assertEqual(response.status_code, 302)
        follow = self.client.get("/places").get_data(as_text=True)
        self.assertIn("Name", follow)

    def test_preview_without_a_file(self):
        response = self.client.post("/places/import/preview", data={})
        self.assertEqual(response.status_code, 302)

    def confirm(self, html, mode, extra_include=()):
        payload = re.search(r'name="payload" value="([^"]*)"', html).group(1)
        import html as html_lib
        include = re.findall(r'name="include" value="(\d+)" checked', html) + list(extra_include)
        return self.client.post("/places/import/confirm", data={
            "payload": html_lib.unescape(payload), "mode": mode, "include": include,
        })

    def test_sales_export_preview_and_confirm_never_expose_sales_data(self):
        original = load_data("places")
        try:
            address_calls = []
            with mock.patch.object(si, "lookup_postcodes", fake_sales_postcodes), \
                    mock.patch.object(si, "lookup_address", lambda q: address_calls.append(q)):
                response = self.client.post(
                    "/places/import/preview",
                    data={"file": (io.BytesIO(sales_workbook()), "CLFSales 12-03-2026.xlsx")},
                    content_type="multipart/form-data",
                )
            html = response.get_data(as_text=True)
            self.assertEqual(response.status_code, 200)
            self.assertEqual(address_calls, [])
            self.assertRegex(html, r"16\s+sales lines")
            self.assertIn("Sells online", html)
            for secret in ("12.34", "INV-77821", "XQ999", "2026-01-15", "Konjac", "07828737806"):
                self.assertNotIn(secret, html)
            self.confirm(html, "merge")
            saved = load_data("places")
            names = [s["name"] for s in saved["shops"]]
            self.assertIn("Natural Health", names)
            self.assertNotIn("Katarzyna Migdal", names)
            self.assertNotIn("The Deli Downstairs", names)
            path = os.path.join(
                os.path.dirname(os.path.dirname(os.path.abspath(__file__))),
                "Emma-Basic-The-Basic-Ingredients", "project", "data", "places.js",
            )
            with open(path, encoding="utf-8") as fh:
                text = fh.read()
            for secret in ("12.34", "INV-77821", "XQ999", "2026-01-15", "Konjac",
                           "07828737806", "Internet Bulk", "Health Food Retailer"):
                self.assertNotIn(secret, text)
        finally:
            _restore(original)

    def test_a_flagged_shop_can_be_ticked_back_in(self):
        original = load_data("places")
        try:
            with mock.patch.object(si, "lookup_postcodes", fake_sales_postcodes):
                html = self.client.post(
                    "/places/import/preview",
                    data={"file": (io.BytesIO(sales_workbook()), "sales.xlsx")},
                    content_type="multipart/form-data",
                ).get_data(as_text=True)
            unticked = re.search(
                r'name="include" value="(\d+)"(?! checked)[^>]*>\s*</td>\s*<td>The Deli Downstairs', html)
            self.assertTrue(unticked)
            self.confirm(html, "merge", extra_include=[unticked.group(1)])
            self.assertIn("The Deli Downstairs", [s["name"] for s in load_data("places")["shops"]])
        finally:
            _restore(original)

    def test_confirm_with_nothing_ticked_saves_nothing(self):
        before = load_data("places")
        html = self.upload([
            ["Name", "Address", "Town", "Postcode"],
            ["Harvey Nichols", "27 Philpot Ln", "Bristol", "BS1 4DJ"],
        ]).get_data(as_text=True)
        payload = re.search(r'name="payload" value="([^"]*)"', html).group(1)
        import html as html_lib
        response = self.client.post("/places/import/confirm", data={
            "payload": html_lib.unescape(payload), "mode": "merge",
        })
        self.assertEqual(response.status_code, 302)
        self.assertEqual(load_data("places"), before)

    def test_confirm_merge_writes_places_js(self):
        original = load_data("places")
        try:
            html = self.upload([
                ["Name", "Address", "Town", "Postcode", "Website"],
                ["Harvey Nichols", "27 Philpot Ln", "Bristol", "BS1 4DJ", "https://hn.example"],
            ]).get_data(as_text=True)
            response = self.confirm(html, "merge")
            self.assertEqual(response.status_code, 302)
            saved = load_data("places")
            self.assertEqual(len(saved["shops"]), len(original["shops"]) + 1)
            added = saved["shops"][-1]
            self.assertEqual(added["name"], "Harvey Nichols")
            self.assertEqual(added["url"], "https://hn.example")
            self.assertEqual((added["lat"], added["lng"]), KNOWN["BS1 4DJ"])
            self.assertEqual(saved["directory"], original["directory"])
            self.assertEqual(saved["featured"], original["featured"])
            path = os.path.join(
                os.path.dirname(os.path.dirname(os.path.abspath(__file__))),
                "Emma-Basic-The-Basic-Ingredients", "project", "data", "places.js",
            )
            with open(path, encoding="utf-8") as fh:
                text = fh.read()
            self.assertTrue(text.startswith("/* Emma Basic — Where to find our products"))
            self.assertIn("window.EB_PLACES = {", text)
            self.assertIn('"name": "Harvey Nichols"', text)
        finally:
            _restore(original)

    def test_confirm_replace_keeps_only_the_file(self):
        original = load_data("places")
        try:
            html = self.upload([
                ["Name", "Address", "Town", "Postcode"],
                ["Harvey Nichols", "27 Philpot Ln", "Bristol", "BS1 4DJ"],
                ["Mackie's", "1 Market St", "Manchester", "M1 1AE"],
            ]).get_data(as_text=True)
            self.confirm(html, "replace")
            saved = load_data("places")
            self.assertEqual([s["name"] for s in saved["shops"]], ["Harvey Nichols", "Mackie's"])
            self.assertEqual(saved["hq"], original["hq"])
        finally:
            _restore(original)

    def test_confirm_with_a_tampered_payload_saves_nothing(self):
        before = load_data("places")
        response = self.client.post("/places/import/confirm", data={"payload": "not json", "mode": "merge"})
        self.assertEqual(response.status_code, 302)
        self.assertEqual(load_data("places"), before)

    def test_hand_editing_still_works_after_an_import(self):
        original = load_data("places")
        try:
            html = self.upload([
                ["Name", "Address", "Town", "Postcode", "Website", "Phone"],
                ["Harvey Nichols", "27 Philpot Ln", "Bristol", "BS1 4DJ", "https://hn.example", "0117 000"],
            ]).get_data(as_text=True)
            self.confirm(html, "merge")
            form = self.client.get("/places").get_data(as_text=True)
            idx = len(original["shops"])
            self.assertIn('name="shop%d_postcode" value="BS1 4DJ"' % idx, form)
            self.assertIn('name="shop%d_url" value="https://hn.example"' % idx, form)
            self.assertIn('name="shop%d_phone" value="0117 000"' % idx, form)
        finally:
            _restore(original)


class ManualSaveTests(unittest.TestCase):
    def setUp(self):
        app.config["TESTING"] = True
        self.client = app.test_client()

    def test_saving_a_shop_without_a_pin_looks_up_its_postcode(self):
        from test_places import _form_from_places
        original = load_data("places")
        try:
            payload = _form_from_places(original)
            idx = len(original.get("shops") or [])
            payload["shop_count"] = str(idx + 1)
            payload["shop%d_name" % idx] = "Harvey Nichols"
            payload["shop%d_city" % idx] = "Bristol"
            payload["shop%d_address" % idx] = "27 Philpot Ln, Bristol"
            payload["shop%d_postcode" % idx] = "bs1 4dj"
            payload["shop%d_url" % idx] = "https://hn.example"
            with mock.patch.object(si, "lookup_postcodes", fake_postcodes):
                response = self.client.post("/places/save", data=payload)
            self.assertEqual(response.status_code, 302)
            added = load_data("places")["shops"][-1]
            self.assertEqual(added["postcode"], "BS1 4DJ")
            self.assertEqual(added["url"], "https://hn.example")
            self.assertEqual((added["lat"], added["lng"]), KNOWN["BS1 4DJ"])
        finally:
            _restore(original)

    def test_older_forms_without_the_new_fields_keep_them(self):
        from test_places import _form_from_places
        original = load_data("places")
        try:
            data = json.loads(json.dumps(original))
            data["shops"][0]["url"] = "https://keep.example"
            data["shops"][0]["postcode"] = "W1A 1AB"
            with app.app_context():
                save_data("places", data)
            payload = _form_from_places(data)
            for key in list(payload):
                if key.endswith(("_url", "_postcode", "_phone")) and key.startswith("shop"):
                    payload.pop(key)
            self.client.post("/places/save", data=payload)
            shop = load_data("places")["shops"][0]
            self.assertEqual(shop["url"], "https://keep.example")
            self.assertEqual(shop["postcode"], "W1A 1AB")
        finally:
            _restore(original)

    def test_lookup_failure_does_not_block_saving(self):
        from test_places import _form_from_places
        original = load_data("places")

        def broken(postcodes):
            raise si.GeocodeError("offline")

        try:
            payload = _form_from_places(original)
            idx = len(original.get("shops") or [])
            payload["shop_count"] = str(idx + 1)
            payload["shop%d_name" % idx] = "Offline Shop"
            payload["shop%d_postcode" % idx] = "BS1 4DJ"
            with mock.patch.object(si, "lookup_postcodes", broken):
                response = self.client.post("/places/save", data=payload)
            self.assertEqual(response.status_code, 302)
            added = load_data("places")["shops"][-1]
            self.assertEqual(added["name"], "Offline Shop")
            self.assertNotIn("lat", added)
        finally:
            _restore(original)


if __name__ == "__main__":
    unittest.main()
